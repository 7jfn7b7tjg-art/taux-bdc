"""Import / export CSV et Excel pour écritures de conversion devise → CAD."""

from __future__ import annotations

import csv
from dataclasses import dataclass
from datetime import date
from decimal import Decimal
from pathlib import Path
from typing import Any, Iterable, Optional, Sequence

from taux_bdc import (
    TauxBdcError,
    TauxChange,
    convertir,
    formater_nombre,
    get_lang,
    journaliser_conversion,
    parse_date,
    parse_montant,
    recuperer_taux,
)

COLONNES_ENTREE = ("date", "devise", "montant", "reference")
COLONNES_SORTIE = (
    "date",
    "devise",
    "montant",
    "reference",
    "date_taux",
    "taux",
    "montant_cad",
    "serie",
    "ajuste",
    "erreur",
)

# Alias d'en-têtes acceptés à l'import (minuscules, sans accents normalisés)
_ALIAS_COLONNES: dict[str, str] = {
    "date": "date",
    "date_transaction": "date",
    "transaction_date": "date",
    "devise": "devise",
    "currency": "devise",
    "ccy": "devise",
    "montant": "montant",
    "amount": "montant",
    "reference": "reference",
    "ref": "reference",
    "memo": "reference",
}


@dataclass
class LigneLot:
    date_tx: Optional[date] = None
    devise: str = ""
    montant: Optional[Decimal] = None
    reference: str = ""
    date_taux: Optional[date] = None
    taux: Optional[Decimal] = None
    montant_cad: Optional[Decimal] = None
    serie: str = ""
    ajuste: bool = False
    erreur: str = ""
    # Valeurs brutes pour ré-export même si parse échoue
    brut_date: str = ""
    brut_devise: str = ""
    brut_montant: str = ""

    def vers_dict_sortie(self) -> dict[str, str]:
        lang = get_lang()
        oui = "oui" if lang == "fr" else "yes"
        non = "non" if lang == "fr" else "no"
        return {
            "date": self.date_tx.isoformat() if self.date_tx else self.brut_date,
            "devise": self.devise or self.brut_devise,
            "montant": (
                formater_nombre(self.montant, 2)
                if self.montant is not None
                else self.brut_montant
            ),
            "reference": self.reference,
            "date_taux": self.date_taux.isoformat() if self.date_taux else "",
            "taux": formater_nombre(self.taux) if self.taux is not None else "",
            "montant_cad": (
                formater_nombre(self.montant_cad, 2) if self.montant_cad is not None else ""
            ),
            "serie": self.serie,
            "ajuste": (oui if self.ajuste else non) if self.taux is not None else "",
            "erreur": self.erreur,
        }


@dataclass
class FicheUnitaire:
    tx: TauxChange
    montant_devise: Decimal
    montant_cad: Decimal
    reference: str = ""

    def vers_ligne_lot(self) -> LigneLot:
        return LigneLot(
            date_tx=self.tx.date_demandee,
            devise=self.tx.devise,
            montant=self.montant_devise,
            reference=self.reference,
            date_taux=self.tx.date_taux,
            taux=self.tx.taux,
            montant_cad=self.montant_cad,
            serie=self.tx.serie,
            ajuste=self.tx.est_ajuste,
        )


def _norm_header(h: str) -> str:
    return (
        h.strip()
        .lower()
        .replace(" ", "_")
        .replace("-", "_")
        .replace("é", "e")
        .replace("è", "e")
        .replace("ê", "e")
    )


def _map_headers(headers: Sequence[str]) -> dict[str, int]:
    mapping: dict[str, int] = {}
    for i, h in enumerate(headers):
        cle = _ALIAS_COLONNES.get(_norm_header(h))
        if cle and cle not in mapping:
            mapping[cle] = i
    manquantes = [c for c in ("date", "devise", "montant") if c not in mapping]
    if manquantes:
        raise TauxBdcError(
            f"Colonnes manquantes / missing columns: {', '.join(manquantes)}. "
            f"Attendu / expected: date, devise, montant [, reference]"
        )
    return mapping


def _cell(row: Sequence[Any], idx: Optional[int]) -> str:
    if idx is None or idx >= len(row):
        return ""
    val = row[idx]
    if val is None:
        return ""
    return str(val).strip()


def _ligne_depuis_cells(
    brut_date: str,
    brut_devise: str,
    brut_montant: str,
    reference: str,
) -> LigneLot:
    ligne = LigneLot(
        brut_date=brut_date,
        brut_devise=brut_devise,
        brut_montant=brut_montant,
        reference=reference,
        devise=brut_devise.strip().upper(),
    )
    try:
        if not brut_date or not brut_devise or not brut_montant:
            raise TauxBdcError("date, devise et montant sont requis.")
        ligne.date_tx = parse_date(brut_date)
        ligne.montant = parse_montant(brut_montant)
        ligne.devise = brut_devise.strip().upper()
    except TauxBdcError as err:
        ligne.erreur = str(err)
    return ligne


def lire_csv(chemin: Path | str) -> list[LigneLot]:
    chemin = Path(chemin)
    with chemin.open("r", encoding="utf-8-sig", newline="") as f:
        sample = f.read(4096)
        f.seek(0)
        try:
            dialect = csv.Sniffer().sniff(sample, delimiters=",;\t")
        except csv.Error:
            dialect = csv.excel
        reader = csv.reader(f, dialect)
        rows = list(reader)
    if not rows:
        return []
    mapping = _map_headers(rows[0])
    lignes: list[LigneLot] = []
    for row in rows[1:]:
        if not any(str(c).strip() for c in row):
            continue
        lignes.append(
            _ligne_depuis_cells(
                _cell(row, mapping.get("date")),
                _cell(row, mapping.get("devise")),
                _cell(row, mapping.get("montant")),
                _cell(row, mapping.get("reference")),
            )
        )
    return lignes


def lire_xlsx(chemin: Path | str) -> list[LigneLot]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise TauxBdcError(
            "openpyxl est requis pour lire les fichiers Excel. "
            "Installez-le : pip install openpyxl"
        ) from exc

    wb = load_workbook(Path(chemin), read_only=True, data_only=True)
    ws = wb.active
    rows_iter = ws.iter_rows(values_only=True)
    try:
        header_row = next(rows_iter)
    except StopIteration:
        wb.close()
        return []
    headers = ["" if c is None else str(c) for c in header_row]
    mapping = _map_headers(headers)
    lignes: list[LigneLot] = []
    for row in rows_iter:
        cells = list(row) if row else []
        if not any(c is not None and str(c).strip() for c in cells):
            continue
        # Dates Excel parfois en datetime
        brut_date = _cell(cells, mapping.get("date"))
        idx_date = mapping.get("date")
        if idx_date is not None and idx_date < len(cells):
            raw = cells[idx_date]
            if hasattr(raw, "date") and callable(raw.date):
                brut_date = raw.date().isoformat()
            elif hasattr(raw, "isoformat") and not isinstance(raw, str):
                brut_date = raw.isoformat()[:10]
        lignes.append(
            _ligne_depuis_cells(
                brut_date,
                _cell(cells, mapping.get("devise")),
                _cell(cells, mapping.get("montant")),
                _cell(cells, mapping.get("reference")),
            )
        )
    wb.close()
    return lignes


def lire_fichier(chemin: Path | str) -> list[LigneLot]:
    chemin = Path(chemin)
    suffix = chemin.suffix.lower()
    if suffix == ".csv":
        return lire_csv(chemin)
    if suffix in (".xlsx", ".xlsm"):
        return lire_xlsx(chemin)
    raise TauxBdcError(f"Format non supporté : {suffix}. Utilisez .csv ou .xlsx.")


def ecrire_csv(chemin: Path | str, lignes: Iterable[LigneLot]) -> None:
    chemin = Path(chemin)
    with chemin.open("w", encoding="utf-8", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=list(COLONNES_SORTIE))
        writer.writeheader()
        for ligne in lignes:
            writer.writerow(ligne.vers_dict_sortie())


def ecrire_xlsx(chemin: Path | str, lignes: Iterable[LigneLot]) -> None:
    try:
        from openpyxl import Workbook
    except ImportError as exc:
        raise TauxBdcError(
            "openpyxl est requis pour écrire les fichiers Excel. "
            "Installez-le : pip install openpyxl"
        ) from exc

    wb = Workbook()
    ws = wb.active
    ws.title = "conversions"
    ws.append(list(COLONNES_SORTIE))
    for ligne in lignes:
        d = ligne.vers_dict_sortie()
        ws.append([d[c] for c in COLONNES_SORTIE])
    wb.save(Path(chemin))


def ecrire_fichier(chemin: Path | str, lignes: Iterable[LigneLot]) -> None:
    chemin = Path(chemin)
    suffix = chemin.suffix.lower()
    if suffix == ".csv":
        ecrire_csv(chemin, lignes)
    elif suffix == ".xlsx":
        ecrire_xlsx(chemin, lignes)
    else:
        raise TauxBdcError(f"Format non supporté : {suffix}. Utilisez .csv ou .xlsx.")


def traiter_ligne(ligne: LigneLot) -> LigneLot:
    """Récupère le taux et convertit une ligne (mute et retourne la même instance)."""
    if ligne.erreur:
        return ligne
    if ligne.date_tx is None or ligne.montant is None or not ligne.devise:
        ligne.erreur = "Ligne incomplète."
        return ligne
    try:
        tx = recuperer_taux(ligne.devise, ligne.date_tx)
        cad = convertir(ligne.montant, tx.taux)
        journaliser_conversion(tx, ligne.montant, cad, reference=ligne.reference)
        ligne.date_taux = tx.date_taux
        ligne.taux = tx.taux
        ligne.montant_cad = cad
        ligne.serie = tx.serie
        ligne.ajuste = tx.est_ajuste
        ligne.erreur = ""
    except TauxBdcError as err:
        ligne.erreur = str(err)
    return ligne


def creer_modele_csv(chemin: Path | str) -> None:
    chemin = Path(chemin)
    with chemin.open("w", encoding="utf-8", newline="") as f:
        writer = csv.writer(f)
        writer.writerow(list(COLONNES_ENTREE))
        writer.writerow(["2026-07-05", "USD", "1500.00", "FAC-123"])
        writer.writerow(["08/07/2026", "EUR", "1000,50", "FAC-124"])


def creer_modele_xlsx(chemin: Path | str) -> None:
    try:
        from openpyxl import Workbook
    except ImportError as exc:
        raise TauxBdcError("openpyxl requis : pip install openpyxl") from exc
    wb = Workbook()
    ws = wb.active
    ws.title = "lot"
    ws.append(list(COLONNES_ENTREE))
    ws.append(["2026-07-05", "USD", "1500.00", "FAC-123"])
    ws.append(["08/07/2026", "EUR", "1000,50", "FAC-124"])
    wb.save(Path(chemin))
